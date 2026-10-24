"""Text extraction for knowledge sources, and chunking that keeps headings with their text."""

from __future__ import annotations

import re
import threading
from pathlib import Path

SUPPORTED = {
    ".pdf", ".docx", ".pptx", ".xlsx", ".md", ".markdown", ".txt", ".text", ".html", ".htm",
    ".csv", ".json", ".rst", ".org", ".tex", ".py", ".js", ".ts", ".tsx", ".jsx", ".rs", ".go",
    ".java", ".c", ".h", ".cpp", ".swift", ".kt", ".rb", ".php", ".sh", ".yaml", ".yml", ".toml",
}
MAX_FILE_BYTES = 60 * 1024 * 1024


_PDFIUM = threading.Lock()  # PDFium is not thread-safe; the indexer reads two files at once


def _pdf(path: Path) -> str:
    """PDFium reads a page in milliseconds; pypdf (pure Python, much slower) stays as the fallback."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return _pdf_slow(path)
    pages: list[str] = []
    try:
        with _PDFIUM:
            doc = pdfium.PdfDocument(str(path))
            try:
                for i in range(len(doc)):
                    page = doc[i]
                    tp = page.get_textpage()
                    text = tp.get_text_range() or ""
                    tp.close()
                    page.close()
                    text = text.replace("\r\n", "\n").replace("\r", "\n")
                    if text.strip():
                        pages.append(f"[page {i + 1}]\n{text}")
            finally:
                doc.close()
    except Exception:  # noqa: BLE001  (damaged or protected file: let pypdf try before giving up)
        return _pdf_slow(path)
    return "\n\n".join(pages)


def _pdf_slow(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[page {i}]\n{text}")
    return "\n\n".join(pages)


def extract(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        out = []
        for p in d.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") and p.text.strip():
                level = re.sub(r"\D", "", style) or "2"
                out.append("#" * min(int(level), 6) + " " + p.text)
            else:
                out.append(p.text)
        for t in d.tables:
            for row in t.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(out)
    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        slides = []
        for i, s in enumerate(prs.slides, 1):
            parts = [f"# Slide {i}"]
            for shape in s.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
                parts.append("Notes: " + s.notes_slide.notes_text_frame.text)
            slides.append("\n".join(parts))
        return "\n\n".join(slides)
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"# Sheet {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    out.append(" | ".join(cells))
        return "\n".join(out)
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(text, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        for h in soup.find_all(re.compile("^h[1-6]$")):
            h.insert_before("\n" + "#" * int(h.name[1]) + " ")
        text = soup.get_text("\n")
    return text


def chunk(text: str, title: str = "", target_words: int = 420, overlap_words: int = 60) -> list[str]:
    """~500-600 tokens per chunk, split on paragraphs, carrying the nearest heading into each chunk."""
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r", "")).strip()
    if not text:
        return []
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    cur: list[str] = []
    words = 0
    heading = ""

    def flush() -> None:
        nonlocal cur, words
        if cur:
            body = "\n\n".join(cur)
            if heading and heading not in body:
                body = heading + "\n\n" + body
            chunks.append(body)
            tail = " ".join(" ".join(cur).split()[-overlap_words:])
            cur = [tail] if overlap_words else []
            words = len(tail.split())

    for p in paras:
        if p.startswith("#"):
            if words > target_words * 0.4:
                flush()
            heading = p.splitlines()[0]
        n = len(p.split())
        if n > target_words * 1.5:  # very long paragraph: hard-split by sentences
            sentences = re.split(r"(?<=[.!?])\s+", p)
            for s in sentences:
                cur.append(s)
                words += len(s.split())
                if words >= target_words:
                    flush()
            continue
        cur.append(p)
        words += n
        if words >= target_words:
            flush()
    if cur and (not chunks or words > overlap_words + 5):
        body = "\n\n".join(cur)
        if heading and heading not in body:
            body = heading + "\n\n" + body
        chunks.append(body)
    return chunks
